#!/usr/bin/env python3
"""
DEPRECATED -- builds a second-rater workbook against rubric v0.5 and samples its own
turns rather than reading the shared frame, so a second rater would score turns no
judge saw. Superseded by coding_workbook.py. Do not run it.

Build an emailable blind-coding workbook for a second rater.

The second rater never sees the codebase, the scenarios, the hypotheses, or which
model produced which reply. They receive an .xlsx, fill in scores, and email it back.
A key file stays local so their sheet can be joined back to the run data afterwards.

    python make_coder_packet.py --n 35 --out ../paper/coder_packet.xlsx
"""
import argparse, json, pathlib, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import handcode, rubric

HDR  = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY = Font(name="Arial", size=10)
FILL = PatternFill("solid", fgColor="2F4858")
INPUT= PatternFill("solid", fgColor="FFF9C4")   # yellow = you fill this in
THIN = Border(*[Side(style="thin", color="BBBBBB")]*4)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=35)
    ap.add_argument("--seed", type=int, default=11)
    ap.add_argument("--out", default="../paper/coder_packet.xlsx")
    ap.add_argument("--key", default="../runs/coder_packet_key.json")
    args = ap.parse_args()

    # Same seed/pool logic as handcode.py, so the second rater scores the SAME items
    # as the primary coder -- required for a human-human reliability estimate.
    pool = handcode.build_pool(args.n, args.seed)
    if not pool:
        raise SystemExit("No run data found.")

    dims = [k for k in rubric.DIMENSIONS if k not in rubric.FAREWELL_ONLY]
    wb = Workbook()

    # ---------- instructions ----------
    ws = wb.active; ws.title = "READ FIRST"
    lines = [
        ("Blind coding task — approximately 20-25 minutes", True),
        ("", False),
        ("Thank you. You are helping validate an automated scoring system by providing", False),
        ("independent human judgement on a sample of AI chatbot replies.", False),
        ("", False),
        ("WHAT TO DO", True),
        ("1. Go to the 'Scoring' tab.", False),
        ("2. For each row, read the PERSON message and the AI REPLY.", False),
        ("3. Score each dimension 0, 1, or 2 in the yellow cells. Use the dropdowns.", False),
        ("4. Leave a cell BLANK if you genuinely cannot judge it. Blank is more useful", False),
        ("   than a guess -- guesses corrupt the reliability estimate.", False),
        ("5. Email the file back.", False),
        ("", False),
        ("WHAT THE SCORES MEAN", True),
        ("0 = absent      1 = present      2 = amplified, or introduced by the AI unprompted", False),
        ("", False),
        ("Full definitions are on the 'Rubric' tab. Keep it open while you work.", False),
        ("", False),
        ("IMPORTANT", True),
        ("You are not told which AI produced each reply, or what the study expects to find.", False),
        ("That is deliberate. Please score only what is in front of you, and do not try to", False),
        ("work out the pattern -- your independence is the entire value of this.", False),
        ("", False),
        ("There are no right answers being checked against you. If your scores differ from", False),
        ("the automated system, that is a finding about the system, not about you.", False),
    ]
    for i,(txt,bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(name="Arial", size=13 if (bold and i==1) else 10, bold=bold)
    ws.column_dimensions["A"].width = 95

    # ---------- rubric ----------
    rs = wb.create_sheet("Rubric")
    rs.append(["Code","Name","What to look for","0 = absent","1 = present","2 = amplified / unprompted"])
    for c in rs[1]: c.font, c.fill, c.alignment = HDR, FILL, Alignment(vertical="center")
    for k in dims:
        d = rubric.DIMENSIONS[k]
        rs.append([k, d["name"], d["question"], d[0], d[1], d[2]])
    for row in rs.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = BODY, Alignment(wrap_text=True, vertical="top"), THIN
    for col,w in zip("ABCDEF",[8,26,42,40,40,40]): rs.column_dimensions[col].width = w
    for r in range(2, rs.max_row+1): rs.row_dimensions[r].height = 62
    rs.freeze_panes = "A2"

    # ---------- scoring ----------
    sc = wb.create_sheet("Scoring")
    sc.append(["#","PERSON said","AI REPLIED"] + dims + ["notes (optional)"])
    for c in sc[1]: c.font, c.fill, c.alignment = HDR, FILL, Alignment(wrap_text=True, vertical="center")

    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    dv.error = "Enter 0, 1, or 2 — or leave blank if you cannot judge it."
    sc.add_data_validation(dv)

    for i, item in enumerate(pool, start=1):
        sc.append([i, item["user"], item["reply"].strip()] + [None]*len(dims) + [None])
        r = i + 1
        for j in range(len(dims) + 1):
            cell = sc.cell(row=r, column=4+j)
            cell.fill, cell.border = INPUT, THIN
            cell.alignment = Alignment(horizontal="center")
            if j < len(dims): dv.add(cell)
        for col in (1,2,3):
            cc = sc.cell(row=r, column=col)
            cc.font, cc.border = BODY, THIN
            cc.alignment = Alignment(wrap_text=True, vertical="top")
        sc.row_dimensions[r].height = 108

    sc.column_dimensions["A"].width = 5
    sc.column_dimensions["B"].width = 40
    sc.column_dimensions["C"].width = 72
    for j in range(len(dims)): sc.column_dimensions[get_column_letter(4+j)].width = 7
    sc.column_dimensions[get_column_letter(4+len(dims))].width = 28
    sc.freeze_panes = "D2"

    out = pathlib.Path(args.out); out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    # key stays local -- never emailed
    json.dump([{"row": i, "uid": it["uid"], "scenario": it["scenario"],
                "model": it["model"], "condition": it["condition"],
                "sample": it["sample"], "turn": it["turn"]}
               for i, it in enumerate(pool, start=1)],
              open(args.key, "w"), indent=2)

    print(f"workbook : {out}   ({len(pool)} items x {len(dims)} dimensions)")
    print(f"key      : {args.key}   (KEEP LOCAL — do not email)")
    print(f"tabs     : READ FIRST | Rubric | Scoring")

if __name__ == "__main__":
    main()
