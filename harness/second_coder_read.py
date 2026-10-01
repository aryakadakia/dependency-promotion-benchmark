#!/usr/bin/env python3
"""
Read the returned second-coder workbook into data/human/second_coded.json.

    python second_coder_read.py ../data/human/second_coder_packet_returned_2026-09-30.xlsx

The workbook carries no scenario, model or condition column, so the answers only
mean anything once they are joined back through second_coder_key.json, which maps
each item number to a frame uid. That join is the one place a silent error would be
unrecoverable -- an off-by-one in the item order would pair every answer with the
wrong reply and still produce a plausible-looking agreement coefficient. So the
exchange text in each shaded row is reconstructed from the frame and compared
byte-for-byte before anything is written, and the per-item dimension lists are
compared against the key. Any mismatch is fatal.

A `skip` is recorded as a skip, not as a 0. The coder was told to use it when a
judgement genuinely could not be made, and collapsing it to "behaviour absent"
would convert an honest abstention into a disagreement with the other coder.
"""
import argparse, json, pathlib, sys
from collections import Counter
from openpyxl import load_workbook
import rubric_v07 as R

ROOT = pathlib.Path(__file__).parent.parent
UID = lambda t: f"{t['scenario']}|{t['model']}|{t['condition']}|{t['sample']}|{t['turn']}"


def blocks_of(ws):
    """One block per item: the shaded exchange row, then one row per question."""
    out, cur = [], None
    for r in range(2, ws.max_row + 1):
        item, code = ws.cell(r, 1).value, ws.cell(r, 2).value
        if item is not None:
            cur = {"item": int(item), "text": ws.cell(r, 3).value, "q": []}
            out.append(cur)
        elif code and cur is not None:
            cur["q"].append({"dim": str(code).strip(),
                             "ans": ws.cell(r, 6).value,
                             "prov": ws.cell(r, 7).value,
                             "note": ws.cell(r, 9).value})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--key", default=str(ROOT / "data" / "human" / "second_coder_key.json"))
    ap.add_argument("--frame", default=str(ROOT / "data" / "frame" / "frame.json"))
    ap.add_argument("--out", default=str(ROOT / "data" / "human" / "second_coded.json"))
    a = ap.parse_args()

    key = {k["item"]: k for k in json.load(open(a.key))}
    turns = {UID(t): t for t in json.load(open(a.frame))["turns"]}
    wb = load_workbook(a.xlsx)
    if wb["CODING"].cell(1, 8).value != "part=S":
        sys.exit("not a second-coder packet (part stamp missing)")
    blocks = blocks_of(wb["CODING"])

    fatal = []
    scores, skipped, notes, blank_prov = {}, [], [], []
    for b in blocks:
        k = key.get(b["item"])
        if k is None:
            fatal.append(f"item {b['item']} is not in the key"); continue
        t = turns.get(k["uid"])
        if t is None:
            fatal.append(f"item {b['item']}: uid {k['uid']} is not in the frame"); continue
        ctx = f"EARLIER:\n{t['prior']}\n\n" if t.get("prior") else ""
        want = f"{ctx}PERSON:\n{t['user']}\n\nASSISTANT REPLY:\n{t['reply']}"
        if (b["text"] or "") != want:
            fatal.append(f"item {b['item']}: the exchange text does not match {k['uid']}")
        if [q["dim"] for q in b["q"]] != k["dims"]:
            fatal.append(f"item {b['item']}: questions {[q['dim'] for q in b['q']]} "
                         f"!= key {k['dims']}")
            continue
        rec = scores.setdefault(k["uid"], {})
        for q in b["q"]:
            d, ans = q["dim"], str(q["ans"]).strip().lower()
            if q["note"]:
                notes.append({"item": b["item"], "uid": k["uid"], "dim": d,
                              "answer": ans, "note": str(q["note"]).strip()})
            if ans == "skip":
                skipped.append({"item": b["item"], "uid": k["uid"], "dim": d}); continue
            if ans not in ("0", "1"):
                fatal.append(f"item {b['item']} {d}: answer {q['ans']!r} is not 1, 0 or skip")
                continue
            rec[d] = int(ans)
            p = None if q["prov"] in (None, "", "—") else str(q["prov"]).strip().lower()
            if R.DIMENSIONS[d].get("provenance") and rec[d] == 1:
                if p in ("person", "assistant"):
                    rec[d + "_init"] = p
                else:
                    # Left blank on a YES. Recorded as a coded 1 with provenance
                    # missing rather than guessed: the presence judgement is the
                    # primary measure and does not depend on it. Kept apart from
                    # the coder's own notes so that counting her notes does not
                    # also count our bookkeeping.
                    blank_prov.append({"item": b["item"], "uid": k["uid"], "dim": d})
            elif p and rec[d] != 1:
                fatal.append(f"item {b['item']} {d}: provenance {p!r} on a {rec[d]}")

    if fatal:
        print("\n".join("FATAL " + f for f in fatal))
        sys.exit(f"{len(fatal)} problem(s); nothing written")

    n_q = sum(len(b["q"]) for b in blocks)
    out = {"source": pathlib.Path(a.xlsx).name, "coder": "second",
           "turns": len(blocks), "questions": n_q,
           "scores": scores, "skipped": skipped, "notes": notes,
           "provenance_blank_on_yes": blank_prov}
    pathlib.Path(a.out).write_text(json.dumps(out, indent=1, sort_keys=True) + "\n")
    vals = [v for s in scores.values() for d, v in s.items() if not d.endswith("_init")]
    print(f"{len(blocks)} turns, {n_q} questions, {len(skipped)} skipped -> {a.out}")
    print(f"coded {len(vals)}: {dict(Counter(vals))}  "
          f"prevalence {sum(vals)/len(vals):.3f}")
    print(f"per dimension: {dict(Counter(d for s in scores.values() for d in s if not d.endswith('_init')))}")
    print(f"{len(notes)} note(s) kept, "
          f"{len(blank_prov)} YES with provenance left blank")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
